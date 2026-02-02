"""
Tests for PPTX exporter.
"""
import os
import pytest
import tempfile
import shutil
from pathlib import Path

from pptx import Presentation

from app.pptx_exporter import PPTXExporter
from app.models import (
    SlideSpec,
    SlideLayout,
    SlideCopy,
    ContentSection,
    DataVizSpec,
    GammaSlideResult,
)


class TestPPTXExporter:
    @pytest.fixture
    def temp_dir(self):
        temp = tempfile.mkdtemp()
        yield temp
        shutil.rmtree(temp)

    @pytest.fixture
    def exporter(self, temp_dir):
        return PPTXExporter(output_dir=temp_dir)

    @pytest.fixture
    def sample_spec(self):
        return SlideSpec(
            slide_id="S01",
            title="Market Overview",
            framework_id="pyramid_executive_summary",
            framework_name="Pyramid / Executive Summary",
            rationale="Best for opening",
            layout=SlideLayout(
                sections=[
                    ContentSection(
                        name="Key Point 1",
                        content_type="text",
                        content="Digital transformation is critical",
                    ),
                    ContentSection(
                        name="Key Point 2",
                        content_type="text",
                        content="Competitors are ahead",
                    ),
                ],
                visual_notes="Use corporate blue",
            ),
            copy=SlideCopy(
                headline="The market has fundamentally shifted",
                body=[
                    "E-commerce penetration accelerated 5 years",
                    "Customer expectations at all-time high",
                    "Competitors investing heavily in digital",
                ],
                footnotes=["Source: Industry Analysis 2024"],
                speaker_notes="Emphasize urgency",
            ),
            data_viz=DataVizSpec(type="none"),
        )

    @pytest.fixture
    def sample_gamma_result(self):
        return GammaSlideResult(
            slide_id="S01",
            success=True,
            gamma_id="gamma_test123",
            structured_content={
                "elements": [
                    {
                        "type": "text",
                        "role": "title",
                        "content": "Market Overview",
                        "position": {"x": 50, "y": 50, "width": 1820, "height": 80},
                        "style": {"fontSize": 36, "fontWeight": "bold"},
                    }
                ]
            },
        )

    def test_create_presentation_from_specs(self, exporter, sample_spec):
        slides_data = [(sample_spec, None)]

        filepath = exporter.create_presentation(
            deck_title="Test Deck",
            slides_data=slides_data,
            author="Test Author",
            subtitle="Test Subtitle",
        )

        assert os.path.exists(filepath)
        assert filepath.endswith(".pptx")

        # Verify the presentation
        prs = Presentation(filepath)
        # Title slide + 1 content slide
        assert len(prs.slides) == 2

    def test_create_presentation_with_gamma_result(self, exporter, sample_spec, sample_gamma_result):
        slides_data = [(sample_spec, sample_gamma_result)]

        filepath = exporter.create_presentation(
            deck_title="Test Deck",
            slides_data=slides_data,
        )

        assert os.path.exists(filepath)
        prs = Presentation(filepath)
        assert len(prs.slides) == 2

    def test_create_presentation_multiple_slides(self, exporter):
        specs = []
        for i in range(5):
            spec = SlideSpec(
                slide_id=f"S{i+1:02d}",
                title=f"Slide {i+1}",
                framework_id="single_message",
                framework_name="Single Message",
                rationale="Test",
                layout=SlideLayout(sections=[]),
                copy=SlideCopy(
                    headline=f"Message {i+1}",
                    body=[f"Point {i+1}"],
                ),
            )
            specs.append((spec, None))

        filepath = exporter.create_presentation(
            deck_title="Multi-Slide Test",
            slides_data=specs,
        )

        prs = Presentation(filepath)
        # Title slide + 5 content slides
        assert len(prs.slides) == 6

    def test_presentation_metadata(self, exporter, sample_spec):
        slides_data = [(sample_spec, None)]

        filepath = exporter.create_presentation(
            deck_title="Metadata Test",
            slides_data=slides_data,
            author="John Doe",
        )

        prs = Presentation(filepath)
        assert prs.core_properties.title == "Metadata Test"
        assert prs.core_properties.author == "John Doe"

    def test_framework_specific_layouts(self, exporter):
        frameworks_to_test = [
            "2x2_matrix",
            "scr",
            "before_after_bridge",
            "pros_cons_tradeoff",
            "funnel",
            "timeline_roadmap",
            "kpi_scorecard",
        ]

        for framework_id in frameworks_to_test:
            spec = SlideSpec(
                slide_id="S01",
                title=f"Test {framework_id}",
                framework_id=framework_id,
                framework_name=framework_id.replace("_", " ").title(),
                rationale="Test",
                layout=SlideLayout(
                    sections=[
                        ContentSection(name=f"Section {i}", content_type="text", content=f"Content {i}")
                        for i in range(4)
                    ]
                ),
                copy=SlideCopy(headline="Test Headline", body=["Point 1", "Point 2"]),
            )

            filepath = exporter.create_presentation(
                deck_title=f"Framework Test: {framework_id}",
                slides_data=[(spec, None)],
            )

            assert os.path.exists(filepath), f"Failed for {framework_id}"

    def test_export_session(self, exporter, sample_spec):
        slides_data = [(sample_spec, None)]

        filepath = exporter.export_session(
            session_id="test_session_123",
            deck_title="Session Export Test",
            slides_data=slides_data,
            author="Exporter",
        )

        assert os.path.exists(filepath)
        # Check session file also created
        session_file = exporter.output_dir / "test_session_123.pptx"
        assert session_file.exists()

    def test_filename_sanitization(self, exporter, sample_spec):
        slides_data = [(sample_spec, None)]

        filepath = exporter.create_presentation(
            deck_title="Test/Deck:With*Special?Chars",
            slides_data=slides_data,
        )

        # Filename should be sanitized
        assert os.path.exists(filepath)
        assert "/" not in os.path.basename(filepath)
        assert ":" not in os.path.basename(filepath)


class TestPPTXExporterSlideContent:
    @pytest.fixture
    def temp_dir(self):
        temp = tempfile.mkdtemp()
        yield temp
        shutil.rmtree(temp)

    @pytest.fixture
    def exporter(self, temp_dir):
        return PPTXExporter(output_dir=temp_dir)

    def test_slide_with_speaker_notes(self, exporter):
        spec = SlideSpec(
            slide_id="S01",
            title="Test",
            framework_id="single_message",
            framework_name="Single Message",
            rationale="Test",
            layout=SlideLayout(sections=[]),
            copy=SlideCopy(
                headline="Main message",
                body=[],
                speaker_notes="Remember to pause here for effect",
            ),
        )

        filepath = exporter.create_presentation(
            deck_title="Notes Test",
            slides_data=[(spec, None)],
        )

        prs = Presentation(filepath)
        # Check the second slide (first is title)
        content_slide = prs.slides[1]
        notes = content_slide.notes_slide.notes_text_frame.text
        assert "pause here for effect" in notes

    def test_slide_with_footnotes(self, exporter):
        spec = SlideSpec(
            slide_id="S01",
            title="Test",
            framework_id="data_chart",
            framework_name="Data Chart",
            rationale="Test",
            layout=SlideLayout(sections=[]),
            copy=SlideCopy(
                headline="Data shows growth",
                body=["Revenue up 20%"],
                footnotes=["Source: Annual Report 2024", "Data as of Q4"],
            ),
        )

        filepath = exporter.create_presentation(
            deck_title="Footnotes Test",
            slides_data=[(spec, None)],
        )

        assert os.path.exists(filepath)

    def test_slide_with_data_viz(self, exporter):
        spec = SlideSpec(
            slide_id="S01",
            title="Data Slide",
            framework_id="data_chart",
            framework_name="Data Chart",
            rationale="Test",
            layout=SlideLayout(sections=[]),
            copy=SlideCopy(headline="Chart Title", body=[]),
            data_viz=DataVizSpec(
                type="bar",
                title="Revenue by Quarter",
                spec={
                    "categories": ["Q1", "Q2", "Q3", "Q4"],
                    "values": [100, 120, 150, 180],
                },
            ),
        )

        filepath = exporter.create_presentation(
            deck_title="Data Viz Test",
            slides_data=[(spec, None)],
        )

        assert os.path.exists(filepath)
